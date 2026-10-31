#!/usr/bin/env python3
"""Export the revised screenplay, eighteen-frame storyboard, viewer and archive."""

from pathlib import Path
from html import escape
import base64
import json
import re
import shutil
import zipfile

from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle
from docx import Document
from docx.shared import Inches, Pt

from pilot_v2_data import TITLE, LOGLINE, DATE, FRAMES, shots, timecode, validate

ROOT = Path(__file__).resolve().parent
SERIES = ROOT.parents[1]
REPO = SERIES.parent
OUTPUT = REPO.parent / 'output' / 'grand-chukuchuku'
BOARD = ROOT / 'Grand-Chukuchuku_Pilot_v2_Cinematic_Storyboard.pdf'
SCRIPT = ROOT / 'Grand-Chukuchuku_Pilot_v2_Screenplay.pdf'
DOCX = SCRIPT.with_suffix('.docx')
VIEWER = ROOT / 'Grand-Chukuchuku_Pilot_v2_Storyboard_Viewer.html'
PACKAGE = SERIES / 'Grand-Chukuchuku_Pilot_v2_Cinematic_Package.zip'
SHEETS = {1: ROOT/'art'/'ACT_I_sheet.png', 2: ROOT/'art'/'ACT_II_sheet.png', 3: ROOT/'art'/'ACT_III_sheet.png'}
W,H = A4
M = 38
BG = colors.HexColor('#101821')
PANEL = colors.HexColor('#192530')
INK = colors.HexColor('#F2EDE3')
MUTED = colors.HexColor('#B2C0C9')
GOLD = colors.HexColor('#E8AE60')
LINE = colors.HexColor('#344452')
TEAL = colors.HexColor('#61BDB2')

# Measured gutter centres in the generated source sheets. Only the PDF/HTML view
# clips the source; original images are preserved without retouching or resampling.
SOURCE_BOXES = {
    1: [(4,4,325,624),(333,4,751,624),(759,4,1147,624),
        (4,631,326,1365),(333,631,751,1365),(759,631,1147,1365)],
    2: [(7,4,368,678),(376,4,746,678),(754,4,1145,678),
        (7,685,368,1362),(376,685,745,1362),(754,685,1145,1362)],
    3: [(4,3,356,669),(366,3,696,669),(706,3,1146,669),
        (4,679,356,1365),(366,679,755,1365),(765,679,1146,1365)],
}

def fonts():
    folder = Path('/usr/share/fonts/truetype/dejavu')
    for name,file in [('Body','DejaVuSans.ttf'),('Bold','DejaVuSans-Bold.ttf'),('Serif','DejaVuSerif.ttf'),('Display','DejaVuSerif-Bold.ttf')]:
        pdfmetrics.registerFont(TTFont(name,str(folder/file)))
    pdfmetrics.registerFontFamily('Body',normal='Body',bold='Bold')

def style(size=9,colour=INK,bold=False,leading=None):
    return ParagraphStyle('p',fontName='Bold' if bold else 'Body',fontSize=size,leading=leading or size*1.4,textColor=colour)

def p(c,text,x,y,width,size=9,colour=INK,bold=False):
    obj=Paragraph(escape(text).replace('\n','<br/>'),style(size,colour,bold))
    _,h=obj.wrap(width,1000)
    if y-h<46: raise ValueError(f'Overflow: {text[:60]} at {y} / {h}')
    obj.drawOn(c,x,y-h)
    return y-h

def label(c,text,x,y,colour=GOLD):
    c.setFillColor(colour);c.setFont('Bold',7.5);c.drawString(x,y,text.upper());return y-16

def chrome(c,page,total,section):
    c.setFillColor(BG);c.rect(0,0,W,H,fill=1,stroke=0)
    c.setStrokeColor(LINE);c.line(M,H-38,W-M,H-38);c.line(M,40,W-M,40)
    c.setFillColor(MUTED);c.setFont('Bold',7);c.drawString(M,H-27,'GRAND-CHUKUCHUKU / PILOT V2')
    c.setFont('Body',6.8);c.drawRightString(W-M,H-27,section.upper())
    c.drawString(M,26,'18 DRAWN KEYFRAMES · 36-SHOT EDIT PLAN · 5:00 TARGET')
    c.drawRightString(W-M,26,f'{page:02d} / {total:02d}')

def box(frame):
    with Image.open(SHEETS[frame['sheet']]) as im: iw,ih=im.size
    left,top,right,bottom=SOURCE_BOXES[frame['sheet']][frame['cell']]
    return (left,top,right,bottom,iw,ih)

def draw_frame(c,frame,x,y,width,height):
    left,top,right,bottom,iw,ih=box(frame)
    scale=min(width/(right-left),height/(bottom-top))
    dw,dh=(right-left)*scale,(bottom-top)*scale
    ax,ay=x+(width-dw)/2,y+(height-dh)/2
    c.setFillColor(colors.HexColor('#080E13'));c.rect(x,y,width,height,fill=1,stroke=0)
    c.saveState();path=c.beginPath();path.rect(ax,ay,dw,dh);c.clipPath(path,stroke=0,fill=0)
    c.drawImage(str(SHEETS[frame['sheet']]),ax-left*scale,ay-(ih-bottom)*scale,width=iw*scale,height=ih*scale,mask='auto')
    c.restoreState()
    c.setStrokeColor(LINE);c.rect(x,y,width,height,fill=0,stroke=1)

def cover(c):
    chrome(c,1,14,'The replacement pilot')
    c.setFillColor(GOLD);c.rect(M,H-80,70,4,fill=1,stroke=0)
    c.setFillColor(INK);c.setFont('Display',30)
    c.drawString(M,H-126,'THE MAN WHO')
    c.drawString(M,H-166,'NEVER RAN')
    p(c,LOGLINE,M,H-191,W-2*M,10.5,MUTED)
    y=155;gap=10;fw=(W-2*M-2*gap)/3
    for n,key in enumerate(['K03','K11','K18']):
        f=next(a for a in FRAMES if a['id']==key)
        draw_frame(c,f,M+n*(fw+gap),y,fw,325)
    label(c,'CHAPTER ONE / FIANGO / CINEMATIC STORYBOARD',M,125)
    p(c,'A real visual sequence: eighteen distinct illustrated moments, a camera and sound plan, and an independent critic review.',M,106,W-2*M,9,MUTED)
    p(c,DATE+' · Author-based dramatic composite · English / Cameroon Pidgin',M,67,W-2*M,7.5,MUTED)
    c.showPage()

def map_page(c):
    chrome(c,2,14,'Story / geography')
    y=H-71
    y=label(c,'THE THRILLER ENGINE',M,y)
    y=p(c,'Who is he when his story cannot save anyone?',M,y,W-2*M,18,INK,True)-13
    y=p(c,'OPEN: an older Zino cannot answer Mother. RETURN: meet him at ten inventing Grand. ESCALATE: later night revision becomes a witness event. CHOOSE: take both friends out. PAY OFF: the storyteller begins an account he cannot rewrite.',M,y,W-2*M,9.4,MUTED)-23
    actrows=[('I / 0:00–1:34','A rule: Grand never runs.'),('II / 1:34–3:24','A crowd makes the rule impossible.'),('III / 3:24–5:00','Zino chooses people over a performance.')]
    for a,b in actrows:
        c.setFillColor(PANEL);c.roundRect(M,y-34,W-2*M,32,4,fill=1,stroke=0)
        p(c,a,M+10,y-6,144,8,GOLD,True);p(c,b,M+165,y-6,330,8.4,INK);y-=40
    y-=12;y=label(c,'SPATIAL CONTRACT / SCREEN DIRECTIONS',M,y)
    # Small schematic is production geography, not a claim about exact coordinates.
    map_y=y-174
    c.setStrokeColor(LINE);c.setFillColor(PANEL);c.roundRect(M,map_y,W-2*M,170,6,fill=1,stroke=1)
    c.setFillColor(TEAL);c.circle(330,map_y+136,15,fill=0,stroke=1)
    p(c,'COMPOUND / MANGO + WELL',M+16,map_y+152,260,8,INK,True)
    c.setStrokeColor(MUTED);c.line(M+12,map_y+98,W-M-12,map_y+98)
    c.line(M+12,map_y+72,W-M-12,map_y+72)
    p(c,'ROAD / CROWD LEFT',M+16,map_y+92,220,8,GOLD,True)
    p(c,'RETURN GATE RIGHT',W-207,map_y+92,162,8,TEAL,True)
    p(c,'PROFESSIONAL SCHOOL / NIGHT REVISION',M+16,map_y+50,W-2*M-30,8,INK,True)
    c.setStrokeColor(TEAL);c.line(315,map_y+35,435,map_y+63);c.line(435,map_y+63,423,map_y+51);c.line(435,map_y+63,419,map_y+63)
    y=map_y-18
    y=p(c,'The wide gate shot comes before obstruction. Return stays toward the already shown gate; no chase, surprise exit or invented perpetrator. Lens numbers describe equivalent framing for animation.',M,y,W-2*M,8.5,MUTED)-18
    y=label(c,'SOURCE AND EDIT LIMITS',M,y)
    p(c,'Age fourteen / approximately 2004 is a proposed composite placement of secondary-school witness memories. Dialogue, individual actions and camera blocking are dramatized. The cards explicitly show the age-ten passage. These are eighteen keyframes, not thirty-six finished drawings or an animated film.',M,y,W-2*M,8.4,MUTED)
    c.showPage()

def frame_spread(c,index):
    first=FRAMES[index*2];second=FRAMES[index*2+1]
    chrome(c,index+3,14,'Drawn sequence / '+first['id']+'–'+second['id'])
    cols=[M,(W/2)+8];width=(W-2*M-16)/2
    for f,x in zip([first,second],cols):
        label(c,f["id"]+' / '+timecode(f['start'])+'–'+timecode(f['end'])+' / SHOTS '+f['shots'],x,H-65)
        p(c,f['title'],x,H-83,width,13,INK,True)
        draw_frame(c,f,x,330,width,385)
        y=311
        y=p(c,f['lens'],x,y,width,7.4,GOLD,True)-8
        y=p(c,f['action'],x,y,width,8.4,INK)-9
        y=label(c,'SOUND / PERFORMANCE',x,y)
        y=p(c,f['sound'],x,y,width,7.8,MUTED)-9
        y=label(c,'CUT / WHY THIS SHOT',x,y)
        p(c,f['cut']+' '+f['purpose'],x,y,width,7.8,MUTED)
    c.showPage()

def shot_pages(c):
    allshots=list(shots())
    for pageindex,subset in enumerate([allshots[:18],allshots[18:]]):
        chrome(c,12+pageindex,14,'36-shot editorial plan')
        p(c,'Every cut has a reason',M,H-67,W-2*M,18,INK,True)
        p(c,'Eighteen illustrated keyframes support this thirty-six-shot target edit. Intermediate movement, inserts and reverse coverage still need animation layout and an animatic.',M,H-98,W-2*M,8,MUTED)
        rows=[['# / FRAME','TIME / CAMERA','ACTION']]
        for s in subset:
            rows.append([f"{s['id']:02d} / {s['frame']}",f"{timecode(s['start'])}–{timecode(s['end'])}\n{s['camera']}",s['action']])
        cells=[[Paragraph(escape(t).replace('\n','<br/>'),style(7.15,INK if i else GOLD,bold=i==0,leading=9.4)) for t in row]for i,row in enumerate(rows)]
        table=Table(cells,colWidths=[57,132,W-2*M-189])
        table.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,-1),.4,LINE),('BACKGROUND',(0,0),(-1,0),PANEL),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5)]))
        _,height=table.wrap(W-2*M,700)
        if height>H-190:raise ValueError(f'Shot table overflow {height}')
        table.drawOn(c,M,H-145-height);c.showPage()

def review_page(c):
    chrome(c,14,14,'Independent review / handoff')
    y=H-71;y=label(c,'QUALITY GATES',M,y)
    y=p(c,'Approval must follow the actual work.',M,y,W-2*M,18,INK,True)-15
    gates=[('01 / THRILLER CAUSALITY','The opening question is answered by a witnessed cause. Zino makes a consequential choice rather than waiting for Mother to reveal a bill.'),('02 / DRAWN SHOTS','Eighteen unique frames show specific actions, shot sizes and viewpoints. Intermediate layout work is listed rather than claimed finished.'),('03 / CONTINUITY','Clear age/time cards; English at home; no age-ten primary-fee confusion; Gideon’s camera stays out of the early passage and closed later.'),('04 / CONSEQUENCE','The adult’s guilt is unknown. No invented evidence, pursuit, rescue or new antagonist. Ordinary friendship and competence precede harm.'),('05 / INDEPENDENT VERDICT','Read CRITIC_FINAL_REVIEW.md for the critic’s final findings on the script and actual sheet images. Remaining limitations belong in the handoff.')]
    for title,text in gates:
        y=label(c,title,M,y)
        y=p(c,text,M,y,W-2*M,9,MUTED)-19
    y=label(c,'NEXT PRODUCTION PASS',M,y)
    y=p(c,'Record the dialogue; draw missing coverage and acting poses; lock the teenage model/wardrobe; verify the witness-event date with the author; then assemble and time a five-minute animatic. No score may turn the violence into triumph. The rejected v1 Artlist music plan does not govern this cut.',M,y,W-2*M,9,MUTED)-18
    y=label(c,'CREDITS / PROVENANCE',M,y)
    p(c,'Story foundations: Nkeh Nkengzincho. Writing, camera and independent critic work are retained in this directory. Visuals: three newly generated storyboard sheets using the author-supplied compound and Grand model references. Source generations and unused first attempts remain preserved. No award or studio affiliation is implied.',M,y,W-2*M,8,MUTED)
    c.showPage()

def build_board():
    c=canvas.Canvas(str(BOARD),pagesize=A4,pageCompression=1)
    c.setTitle('Grand-Chukuchuku / Pilot v2 / '+TITLE);c.setAuthor('Based on the life of Nkeh Nkengzincho')
    cover(c);map_page(c)
    for i in range(9):frame_spread(c,i)
    shot_pages(c);review_page(c);c.save()

def markdown_script():
    base=(ROOT/'SCRIPT_SCENES.md').read_text()
    lines=[base,'\n## Canonical 36-shot editorial plan\n','| Shot | Time | Keyframe | Camera | Action |','| --- | --- | --- | --- | --- |']
    for s in shots():lines.append(f"| {s['id']:02d} | {timecode(s['start'])}–{timecode(s['end'])} | {s['frame']} | {s['camera']} | {s['action']} |")
    path=SCRIPT.with_suffix('.md');path.write_text('\n'.join(lines));return base

def script_chrome(c,doc):
    c.setFillColor(colors.HexColor('#FAF8F2'));c.rect(0,0,W,H,fill=1,stroke=0)
    c.setFillColor(colors.HexColor('#263E46'));c.setFont('Body',7)
    c.drawString(M,H-30,'GRAND-CHUKUCHUKU / PILOT V2 / '+TITLE)
    c.drawString(M,28,DATE+' / Proposed composite / '+str(doc.page))

def build_script(base):
    flow=[];doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.7)
    doc.styles['Normal'].font.size=Pt(10.5)
    dark=colors.HexColor('#203B42')
    for raw in base.splitlines():
        line=raw.strip()
        if not line:flow.append(Spacer(1,5));continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'));title=line.lstrip('# ').strip()
            if level==2 and title.startswith('ACT '):flow.append(PageBreak());doc.add_page_break()
            font='Display' if level<3 else 'Bold';size=21 if level==1 else 13 if level==2 else 10.5
            st=ParagraphStyle('heading',fontName=font,fontSize=size,leading=size*1.3,textColor=dark,spaceBefore=8,spaceAfter=8,keepWithNext=True)
            flow.append(Paragraph(escape(title),st));doc.add_heading(title,min(level-1,2))
        else:
            rich=re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',escape(line))
            st=ParagraphStyle('body',fontName='Body',fontSize=9.3,leading=13.4,textColor=dark,spaceAfter=5)
            flow.append(Paragraph(rich,st));doc.add_paragraph(re.sub(r'\*\*(.*?)\*\*',r'\1',line))
    SimpleDocTemplate(str(SCRIPT),pagesize=A4,leftMargin=M,rightMargin=M,topMargin=54,bottomMargin=50,title=TITLE,author='Based on the life of Nkeh Nkengzincho').build(flow,onFirstPage=script_chrome,onLaterPages=script_chrome)
    doc.core_properties.title=TITLE;doc.core_properties.author='Based on the life of Nkeh Nkengzincho';doc.save(DOCX)

def build_viewer():
    encoded={str(n):'data:image/png;base64,'+base64.b64encode(path.read_bytes()).decode()for n,path in SHEETS.items()}
    data=[]
    for f in FRAMES:
        v=dict(f);left,top,right,bottom,iw,ih=box(f)
        v['crop']=[left/iw,top/ih,(right-left)/iw,(bottom-top)/ih];data.append(v)
    template='''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Grand-Chukuchuku / Pilot v2</title>
<style>*{box-sizing:border-box}body{margin:0;background:#101821;color:#f2ede3;font:16px/1.5 system-ui}header,main,footer{max-width:1160px;margin:auto;padding:24px}header{border-bottom:1px solid #344452}small,.muted{color:#b2c0c9}h1{font:clamp(24px,4vw,46px)/1.1 Georgia;margin:.4em 0}h2{font-size:26px}button{background:#253848;color:#f2ede3;border:1px solid #536778;border-radius:6px;padding:12px;cursor:pointer}button:focus-visible,input:focus-visible{outline:3px solid #e8ae60;outline-offset:3px}.stage{display:grid;grid-template-columns:minmax(240px,46%) 1fr;gap:32px}.art{height:min(68vh,700px);overflow:hidden;position:relative;background:#080e13}.art img{position:absolute;max-width:none}label{display:block;color:#e8ae60;font-size:12px;letter-spacing:1px;margin-top:20px}p{margin:.4em 0 1em}nav{display:flex;gap:12px;align-items:center;margin-bottom:20px}input{width:100%;accent-color:#e8ae60}#chips{display:flex;gap:8px;flex-wrap:wrap;margin-top:18px}.chip{min-width:46px;padding:8px}.chip[aria-current=true]{background:#e8ae60;color:#101821}.warning{border-left:3px solid #e8ae60;padding:12px 16px;background:#192530}a{color:#e8ae60}@media(max-width:650px){.stage{grid-template-columns:1fr}.art{height:75vh}header,main,footer{padding:18px}}</style>
<header><small>GRAND-CHUKUCHUKU / CHAPTER ONE / PILOT V2</small><h1>THE MAN WHO NEVER RAN</h1><p>18 drawn keyframes · 36-shot editorial plan · five-minute target</p></header>
<main><nav aria-label="Storyboard navigation"><button id="prev" aria-label="Previous keyframe">← Previous</button><span id="count"></span><button id="next" aria-label="Next keyframe">Next →</button></nav><label for="scrub">JUMP THROUGH THE FIVE-MINUTE EDIT</label><input id="scrub" type="range" min="0" max="299" value="0" aria-label="Pilot time in seconds"><p id="time" class="muted"></p>
<div class="stage"><div class="art" id="art"><img id="img" alt=""></div><section aria-live="polite"><small id="meta"></small><h2 id="title"></h2><label>CAMERA</label><p id="lens"></p><label>ACTION</label><p id="action"></p><label>SOUND / PERFORMANCE</label><p id="sound"></p><label>CUT</label><p id="cut"></p><label>WHY THIS SHOT</label><p id="purpose"></p></section></div><div id="chips" aria-label="Choose a keyframe"></div></main>
<footer><p class="warning">This is a still storyboard viewer, not a timed animated performance. Approximate age fourteen / 2004 is proposed placement of a secondary-school witness composite. Dialogue, route and individual actions are dramatized; accusation is never proof of guilt.</p><p class="muted">Use ← / → to compare frames. The screenplay, independent critic review and original sheet images accompany this file.</p></footer>
<script>const frames=__FRAMES__;const sheets=__SHEETS__;let current=0;const el=id=>document.getElementById(id);const tc=s=>`${Math.floor(s/60)}:${String(s%60).padStart(2,'0')}`;function crop(){const f=frames[current],box=el('art'),img=el('img'),r=f.crop;const native=img.naturalWidth;if(!native)return;const nw=img.naturalWidth,nh=img.naturalHeight;const cw=r[2]*nw,ch=r[3]*nh;const scale=Math.min(box.clientWidth/cw,box.clientHeight/ch);img.style.width=nw*scale+'px';img.style.height=nh*scale+'px';img.style.left=(box.clientWidth-cw*scale)/2-r[0]*nw*scale+'px';img.style.top=(box.clientHeight-ch*scale)/2-r[1]*nh*scale+'px';}function render(i){current=Math.max(0,Math.min(17,i));const f=frames[current];['title','lens','action','sound','cut','purpose'].forEach(k=>el(k).textContent=f[k]);el('count').textContent=`${current+1} / 18`;el('meta').textContent=`${f.id} / SHOTS ${f.shots} / ${tc(f.start)}–${tc(f.end)}`;el('time').textContent=`${tc(f.start)}–${tc(f.end)}`;el('scrub').value=f.start;el('img').alt=f.title+'. '+f.action;el('img').onload=crop;el('img').src=sheets[String(f.sheet)];crop();el('prev').disabled=current===0;el('next').disabled=current===17;document.querySelectorAll('.chip').forEach((b,j)=>b.setAttribute('aria-current',j===current));}frames.forEach((f,i)=>{const b=document.createElement('button');b.className='chip';b.textContent=f.id;b.onclick=()=>render(i);el('chips').append(b)});el('prev').onclick=()=>render(current-1);el('next').onclick=()=>render(current+1);el('scrub').oninput=e=>{const s=Number(e.target.value);render(frames.findIndex(f=>s>=f.start&&s<f.end));};document.onkeydown=e=>{if(e.target.tagName==='INPUT')return;if(e.key==='ArrowRight')render(current+1);if(e.key==='ArrowLeft')render(current-1);};window.onresize=crop;render(0);</script></html>'''
    VIEWER.write_text(template.replace('__FRAMES__',json.dumps(data,ensure_ascii=False)).replace('__SHEETS__',json.dumps(encoded)))

def package():
    includes=[BOARD,SCRIPT,DOCX,SCRIPT.with_suffix('.md'),VIEWER,ROOT/'README.md',ROOT/'pilot_v2_data.py',Path(__file__),ROOT/'SCRIPT_SCENES.md']+list(SHEETS.values())
    includes+=sorted(ROOT.glob('CRITIC*.md'))+[ROOT/'WRITER_THRILLER_TREATMENT.md',ROOT/'DIRECTOR_CAMERA_BRIEF.md']
    includes+=[SERIES/'CONTINUITY_LOG.md',SERIES/'visual-reference'/'VISUAL_BASELINE_V1.md']
    includes+=[SERIES/'writers-room'/'CONTEXT_REVIEW_2026-10-30.md',SERIES/'writers-room'/'LATEST_CHARACTER_CONTINUITY_REVIEW_2026-10-30.md',SERIES/'writers-room'/'NIGHT_WARDROBE_AND_TURNAROUNDS_REVIEW_2026-10-30.md',SERIES/'character-bible'/'Fiango_Life_and_Characters_Update_2026-10-30.md',SERIES/'character-bible'/'Fiango_Life_and_Characters_Update_2026-10-30.pdf']
    includes+=[SERIES/'writers-room'/'LIZZY_MIRABEL_DISTANCE_REVIEW_2026-10-30.md']
    includes+=[SERIES/'writers-room'/'COMPLEXION_AND_UNIFORM_REVIEW_2026-10-31.md']
    with zipfile.ZipFile(PACKAGE,'w',zipfile.ZIP_DEFLATED)as z:
        for file in includes:
            if file.exists():
                arc=file.relative_to(ROOT)if file.is_relative_to(ROOT)else Path('Source_Notes')/file.name
                z.write(file,arc)
    OUTPUT.mkdir(parents=True,exist_ok=True)
    for file in [BOARD,SCRIPT,PACKAGE]:shutil.copy2(file,OUTPUT/file.name)

def main():
    validate();fonts()
    for path in SHEETS.values():assert path.is_file()and path.stat().st_size>1000,path
    base=markdown_script();build_board();build_script(base);build_viewer();package()
    print(json.dumps({'board':str(BOARD),'script':str(SCRIPT),'viewer':str(VIEWER),'package':str(PACKAGE),'frames':18,'shots':36,'target_seconds':300},indent=2))

if __name__=='__main__':main()
