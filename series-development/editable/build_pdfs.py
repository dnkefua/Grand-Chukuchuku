#!/usr/bin/env python3
"""Build the series workbook, standalone chapters, editable text, and archive."""
from pathlib import Path
from html import escape
import json
import re
import zipfile

from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, Table, TableStyle

ROOT = Path(__file__).resolve().parent.parent
DATA = json.loads((ROOT / "editable" / "series_outline.json").read_text())
W, H = A4
M = 45
CW = W - 2 * M
INK = colors.HexColor("#183C39")
MUTED = colors.HexColor("#556763")
PAPER = colors.HexColor("#FAF8F2")
GOLD = colors.HexColor("#BD8D42")
PALE = colors.HexColor("#E8EEE7")
LINE = colors.HexColor("#D2DAD2")
WHITE = colors.white

FONT_ROOT = Path("/usr/share/fonts/truetype/dejavu")
for name, file in [
    ("Body", "DejaVuSans.ttf"), ("BodyB", "DejaVuSans-Bold.ttf"),
    ("BodyI", "DejaVuSans-Oblique.ttf"), ("Title", "DejaVuSerif.ttf"),
    ("TitleB", "DejaVuSerif-Bold.ttf"),
]:
    pdfmetrics.registerFont(TTFont(name, str(FONT_ROOT / file)))
pdfmetrics.registerFontFamily("Body", normal="Body", bold="BodyB", italic="BodyI", boldItalic="BodyB")

STYLES = {
    "body": ParagraphStyle("body", fontName="Body", fontSize=10, leading=14.8, textColor=INK, alignment=TA_LEFT),
    "small": ParagraphStyle("small", fontName="Body", fontSize=8.7, leading=12.5, textColor=MUTED),
    "title": ParagraphStyle("title", fontName="TitleB", fontSize=26, leading=31, textColor=INK),
    "sub": ParagraphStyle("sub", fontName="BodyB", fontSize=11.3, leading=16, textColor=INK),
    "deck": ParagraphStyle("deck", fontName="Body", fontSize=10.3, leading=15, textColor=MUTED),
    "quote": ParagraphStyle("quote", fontName="Title", fontSize=14.2, leading=20, textColor=INK),
    "cell": ParagraphStyle("cell", fontName="Body", fontSize=8.7, leading=12.1, textColor=INK),
    "cellB": ParagraphStyle("cellB", fontName="BodyB", fontSize=8.7, leading=12.1, textColor=INK),
}

def rich(text):
    return escape(text).replace("\n", "<br/>")

class Layout:
    def __init__(self, c, page, total, section):
        self.c, self.page, self.total, self.section = c, page, total, section
        c.setFillColor(PAPER)
        c.rect(0, 0, W, H, fill=1, stroke=0)
        c.setFillColor(MUTED)
        c.setFont("BodyB", 7.6)
        c.drawString(M, H - 32, "GRAND-CHUKUCHUKU")
        c.setFont("Body", 7.4)
        c.drawRightString(W - M, H - 32, section.upper())
        c.setStrokeColor(LINE)
        c.line(M, H - 43, W - M, H - 43)
        c.line(M, 46, W - M, 46)
        c.setFont("Body", 7.2)
        c.setFillColor(MUTED)
        c.drawString(M, 31, "AUTHOR-BASED DEVELOPMENT DRAFT  |  28 OCT 2026")
        c.setFont("BodyB", 8)
        c.drawRightString(W - M, 31, f"{page:02d} / {total:02d}")
        self.y = H - 67

    def p(self, text, style="body", gap=8, x=None, width=None):
        x = M if x is None else x
        width = CW if width is None else width
        obj = Paragraph(rich(text), STYLES[style])
        _, height = obj.wrap(width, H)
        if self.y - height < 62:
            raise ValueError(f"Overflow page {self.page} ({self.section}): {text[:85]} at {self.y:.1f}, height {height:.1f}")
        obj.drawOn(self.c, x, self.y - height)
        self.y -= height + gap

    def label(self, text):
        self.c.setFillColor(GOLD)
        self.c.setFont("BodyB", 8)
        self.c.drawString(M, self.y - 8, text.upper())
        self.y -= 24

    def title(self, label, title, deck=None):
        self.label(label)
        self.p(title, "title", 12)
        if deck:
            self.p(deck, "deck", 13)

    def sub(self, title, text=None):
        self.p(title, "sub", 5)
        if text:
            self.p(text)

    def note(self, label, text):
        content = Paragraph("<b>" + rich(label) + "</b> " + rich(text), STYLES["small"])
        _, height = content.wrap(CW - 24, H)
        box_h = height + 20
        if self.y - box_h < 62:
            raise ValueError(f"Note overflow on {self.page}: {label}")
        self.c.setFillColor(PALE)
        self.c.roundRect(M, self.y - box_h, CW, box_h, 6, stroke=0, fill=1)
        content.drawOn(self.c, M + 12, self.y - height - 10)
        self.y -= box_h + 12

    def table(self, rows, widths, bold_first=True, size="cell", pad=6):
        parsed = []
        for row in rows:
            parsed.append([Paragraph(rich(str(cell)), STYLES["cellB" if j == 0 and bold_first else size]) for j, cell in enumerate(row)])
        obj = Table(parsed, colWidths=widths, hAlign="LEFT")
        obj.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), pad),
            ("RIGHTPADDING", (0, 0), (-1, -1), pad),
            ("TOPPADDING", (0, 0), (-1, -1), pad),
            ("BOTTOMPADDING", (0, 0), (-1, -1), pad),
            ("ROWBACKGROUNDS", (0, 0), (-1, -1), [WHITE, PALE]),
            ("LINEBELOW", (0, 0), (-1, -1), 0.35, LINE),
        ]))
        _, height = obj.wrap(CW, H)
        if self.y - height < 62:
            raise ValueError(f"Table overflow page {self.page}: {height} available {self.y - 62}")
        obj.drawOn(self.c, M, self.y - height)
        self.y -= height + 12

    def rules(self, count=2):
        self.c.setStrokeColor(LINE)
        self.c.setLineWidth(0.45)
        for _ in range(count):
            if self.y - 13 < 62:
                raise ValueError(f"Writing rules overflow page {self.page}")
            self.c.line(M, self.y - 5, W - M, self.y - 5)
            self.y -= 13
        self.y -= 4

def page_end(c):
    c.showPage()

def bookmark(c, key, title):
    c.bookmarkPage(key)
    c.addOutlineEntry(title, key, level=0)

def cover(c):
    c.setFillColor(PAPER)
    c.rect(0, 0, W, H, fill=1, stroke=0)
    c.setFillColor(GOLD)
    c.rect(M, H - 87, 44, 4, fill=1, stroke=0)
    c.setFont("BodyB", 9)
    c.setFillColor(MUTED)
    c.drawString(M, H - 61, "AN AUTOBIOGRAPHICAL ANIMATED THRILLER")
    c.setFillColor(INK)
    c.setFont("TitleB", 43)
    c.drawString(M - 1, H - 160, "GRAND-")
    c.setFont("TitleB", 38)
    c.drawString(M - 1, H - 210, "CHUKUCHUKU")
    c.setFont("Body", 12)
    c.drawString(M, H - 250, "Series synopsis & five-chapter development workbook")
    c.setStrokeColor(LINE)
    c.line(M, H - 272, W - M, H - 272)
    para = Paragraph(rich("A boy invents an impossible hero. Life asks him to become a man his mother can depend on."), STYLES["quote"])
    _, h = para.wrap(CW - 15, 150)
    para.drawOn(c, M, H - 304 - h)
    c.setFillColor(GOLD)
    c.setFont("BodyB", 8.5)
    c.drawString(M, 392, "READING MAP")
    entries = [
        ("Series foundation & synopsis", "02-03"),
        ("People, world & chronology", "04-06"),
        ("01  Fiango", "07-08"),
        ("02  The Price of Leaving", "09-10"),
        ("03  America on Borrowed Time", "11-12"),
        ("04  The Betrayal", "13-14"),
        ("05  Rebuilding", "15-16"),
        ("Expansion guide", "17"),
    ]
    y = 366
    c.setFont("Body", 9.2)
    for title, pages in entries:
        c.setFillColor(INK)
        c.drawString(M, y, title)
        c.setFillColor(MUTED)
        c.drawRightString(321, y, pages)
        y -= 24
    # Abstract compound motif: symbolic, not a reconstruction of the real home.
    center_x, center_y = 438, 283
    c.setStrokeColor(LINE)
    c.setLineWidth(1)
    for radius in [35, 60, 88]:
        c.circle(center_x, center_y, radius, stroke=1, fill=0)
    import math
    for a in [0, 52, 104, 156, 208, 260, 312]:
        x = center_x + 88 * math.cos(math.radians(a))
        y = center_y + 88 * math.sin(math.radians(a))
        c.setFillColor(INK)
        c.roundRect(x - 6, y - 6, 12, 12, 2, fill=1, stroke=0)
    c.setFillColor(GOLD)
    c.circle(center_x, center_y, 13, stroke=0, fill=1)
    c.setFont("Body", 7.1)
    c.setFillColor(MUTED)
    c.drawCentredString(center_x, 172, "HOME / IMAGINATION / RETURN")
    c.setStrokeColor(LINE)
    c.line(M, 117, W - M, 117)
    c.setFillColor(INK)
    c.setFont("BodyB", 10)
    c.drawString(M, 94, "Based on the life of Nkeh Nkengzincho (Zino)")
    c.setFont("Body", 8.5)
    c.setFillColor(MUTED)
    c.drawString(M, 73, "Development draft  |  28 October 2026  |  For author expansion")
    c.drawString(M, 52, "Includes proposed dramatic structure; not a completed screenplay or verified biography.")
    page_end(c)

def foundation(c):
    l = Layout(c, 2, 17, "Series foundation")
    bookmark(c, "foundation", "Series foundation")
    l.title("The series", "An unfinished promise", "Suspense, heartbreak, and hope, grounded in one family's experience.")
    l.sub("Logline", DATA["logline"])
    l.sub("The premise", DATA["premise"])
    l.table([
        ["Format", "Animated thriller; episodes of 1-3 minutes; vertical 9:16 intended."],
        ["Audience & language", "Cameroonians, Africans across the continent, and the diaspora. English and Cameroon Pidgin; English subtitles proposed."],
        ["Tone & look", "Hard-R material: direct depictions of violence and adult consequences, balanced by humor rooted in character and daily life. Original expressive 2D animation, with The Boondocks as a quality reference."],
        ["Production", "Author-supplied voices; AI-assisted visuals. Build a hard-R master; later create separate platform-release cuts where TikTok rules require them. Budget and release schedule remain to be decided."],
    ], [110, CW - 110])
    l.sub("The engine of the series", "Each stage tests the same belief: if Zino is clever and works hard enough, can he keep his mother safe and make her proud? His talent creates opportunities; temper, people-pleasing, and uncertain self-awareness complicate what follows.")
    l.sub("The whole-series arc", "Act I: Fiango and the cost of leaving establish the dream. Act II: America and the alleged betrayal test it. Act III: rebuilding asks what success and happiness actually mean. Chapters are story movements, not fixed seasons; their episode counts remain open.")
    l.note("Scope", "Biography is drawn from Zino's account in this conversation. Act structures, episode seeds, emotional transitions, and future scene arrangements are proposals for expansion, not additional factual claims.")
    page_end(c)

def synopsis(c):
    l = Layout(c, 3, 17, "Unified synopsis")
    bookmark(c, "synopsis", "Unified series synopsis")
    l.title("The complete story", "From Fiango to Dubai", "One continuous synopsis connecting all five chapters.")
    for p in DATA["synopsis"]:
        l.p(p, gap=10)
    page_end(c)

def people(c):
    l = Layout(c, 4, 17, "People & relationships")
    bookmark(c, "people", "People and relationships")
    l.title("Character foundation", "The people around Zino", "Develop each person beyond the role they play in his success or suffering.")
    for name, text in DATA["characters"][:7]:
        l.p(name, "sub", 3)
        l.p(text, "small", 9)
    l.note("Language", "Zino uses Pidgin with friends and English at home; French and English feature in school. Let remembered speech guide individual voices. Do not make every character sound alike.")
    page_end(c)

def world(c):
    l = Layout(c, 5, 17, "Fiango world")
    bookmark(c, "world", "Fiango world")
    l.title("Fiango", "The world that made him", "The compound, school, and street life should carry as much story weight as plot.")
    for name, text in DATA["world"]:
        l.p(name, "sub", 3)
        l.p(text, "small", 8)
    l.sub("Later relationships and story device")
    for name, text in DATA["characters"][7:]:
        l.p(name, "sub", 2)
        l.p(text, "small", 6)
    l.note("Hard-R depiction rule", "Show violence and adult consequences directly when the story requires it. In jungle-justice scenes, keep the accused person's guilt unproven and follow Zino's witness perspective, physical fear, and aftermath. The brutality must have narrative consequence.")
    page_end(c)

def chronology(c):
    l = Layout(c, 6, 17, "Chronology & open details")
    bookmark(c, "timeline", "Chronology and open details")
    l.title("Working chronology", "What belongs where", "Anchors supplied by Zino; exact months and several connecting details remain open.")
    l.table(DATA["timeline"], [104, CW - 104], pad=5)
    l.sub("Public history to verify", "Place the Ambazonian crisis, Cameroon football victories, and the Biya, Obama, and Trump presidencies against the confirmed personal timeline. Identify the intended AFCON victory. Clarify whether the World Trade Center reference means the 1993 bombing or the September 11, 2001 attacks. Public figures are context; no specific policy has yet been established as causing Zino's experience.")
    l.note("Continuity rule", "Keep nursing licensing, university payments, immigration status, and business ownership as distinct issues. Date Zino's father's death, the aunt's killing, family displacement, and the reported jungle-justice incidents before deciding where to dramatize them. Never turn a proposed scene into an established memory simply by repeating it.")
    page_end(c)

def chapter_outline(c, ch, page, total):
    l = Layout(c, page, total, f"Chapter {ch['number']} / Three-act outline")
    bookmark(c, f"chapter{ch['number']}", f"{ch['number']:02d}  {ch['title']}")
    l.title(f"Chapter {ch['number']:02d}  |  {ch['period']}", ch["title"], ch["focus"])
    l.note("Dramatic question", ch["question"])
    for act in ch["acts"]:
        l.p(act["title"], "sub", 4)
        l.p(act["text"], "body", 4)
        l.p("Turning point: " + act["turn"], "small", 11)
    page_end(c)

def chapter_workbook(c, ch, page, total):
    l = Layout(c, page, total, f"Chapter {ch['number']} / Expansion workbook")
    l.title(f"Chapter {ch['number']:02d}  |  Development page", "Make this chapter yours", ch["title"])
    l.table([["At the beginning", ch["start"]], ["At the end", ch["end"]]], [105, CW - 105])
    l.sub("Three possible short episodes")
    for i, (title, text) in enumerate(ch["episodes"], 1):
        l.p(f"{i}. {title}", "sub", 2)
        l.p(text, "small", 7)
    l.sub("Memory prompts")
    for i, q in enumerate(ch["prompts"], 1):
        l.p(f"{i}. {q}", "small", 3)
        l.rules(1)
    l.note("Grand-Chukuchuku motif", ch["device"])
    l.p("Details still open: " + ch["continuity"], "small", 0)
    page_end(c)

def expansion(c):
    l = Layout(c, 17, 17, "Expansion guide")
    bookmark(c, "expansion", "How to expand the chapters")
    l.title("Your next writing pass", "From outline to lived scenes", "Use the chapter PDFs independently; the editable text contains the same structure.")
    l.sub("1. Capture one memory at a time", "Record the year or approximate age, location, people present, what Zino wanted, what happened, and what changed. Add sensory details: food, church, football, school routines, music, slang, discipline, celebrations, and neighborhood sounds. Separate exact remembered words from reconstructed dialogue.")
    l.sub("2. Identify the knowledge behind it", "Label material as personal recollection, something another person reported, a document or public record, or a proposed dramatization. Keep private evidence and real identities in your own notes. The outline does not independently verify allegations or identify responsibility for the aunt's death.")
    l.sub("3. Give each short episode a complete movement", "Act I: a concrete desire and an immediate problem. Act II: an attempt, complication, or difficult choice. Act III: a consequence that resolves or changes the immediate problem and opens the next question. A working two-minute rhythm is 0-20 seconds, 20-95 seconds, and 95-120 seconds; adjust for performance and pacing.")
    l.sub("4. Make the ending earn the next episode", "Use a discovery, changed relationship, practical deadline, or consequential choice. Pay off the previous episode's question. The football humor and fantasy sequences should deepen Zino's character, while silence and ordinary actions can carry grief without spectacle.")
    l.sub("5. Preserve the truthful destination", "Choose a real present-day decision or relationship change to support the proposed redemption arc. Do not finish the mother's house in the story unless it has actually been completed, or present hopes about venture value as realized money. Exact episode counts and season breaks should follow the expanded material.")
    l.note("For the first expansion", "Start with Chapter 1: one real football afternoon, one school-fee exclusion, one meal after a boarding-school escape, and one conversation with your mother. These memories can establish the show's voice before the later thriller develops.")
    l.p("Working scene card", "sub", 5)
    l.p("Date / place / people: _________________________________\nWhat I wanted: _______________________________________\nWhat stood in the way: __________________________________\nWhat I did, and what changed: ____________________________\nSource or uncertainty: __________________________________", "small", 0)
    page_end(c)

def make_canvas(path):
    c = canvas.Canvas(str(path), pagesize=A4, pageCompression=1)
    c.setTitle("Grand-Chukuchuku | Series Synopsis & Chapter Workbook")
    c.setAuthor("Based on the life of Nkeh Nkengzincho (Zino)")
    c.setSubject("Autobiographical animated thriller: unified synopsis and five three-act chapter outlines")
    return c

def export_markdown():
    lines = ["# GRAND-CHUKUCHUKU", "", "Series synopsis and five-chapter development workbook", "", f"Based on the life of {DATA['author']}", "", f"Development draft: {DATA['date']}", "", "Biography below is drawn from Zino's account. Structure, episode seeds, and emotional transitions are proposed for expansion, not additional biographical claims.", "", "## Logline", "", DATA["logline"], "", "## Unified synopsis", ""]
    lines[10:10] = ["## Series foundation", "", DATA["premise"], "", "**Format:** Hard-R animated thriller; 1-3 minute episodes; vertical 9:16 intended. English and Cameroon Pidgin, with English subtitles proposed.", "", "**Audience:** Cameroonians, Africans across the continent, and the diaspora.", "", "**Production:** Author-supplied voices; AI-assisted visuals. Original expressive 2D animation, with The Boondocks as a quality reference. Build a hard-R master, then make separate platform-release cuts where TikTok rules require them. Budget and release schedule remain open.", "", "**Overall arc:** Act I establishes the dream through Fiango and the cost of leaving. Act II tests it in America and the alleged betrayal. Act III asks what rebuilding, success, and happiness mean.", ""]
    for p in DATA["synopsis"]:
        lines += [p, ""]
    lines += ["## People", ""]
    for name, text in DATA["characters"]:
        lines += [f"**{name}:** {text}", ""]
    lines += ["## Fiango world", ""]
    for name, text in DATA["world"]:
        lines += [f"**{name}:** {text}", ""]
    lines += ["## Chronology", ""]
    for period, text in DATA["timeline"]:
        lines += [f"- **{period}:** {text}"]
    for ch in DATA["chapters"]:
        part = [f"# Chapter {ch['number']}: {ch['title']}", "", f"**Period:** {ch['period']}", "", f"**Focus:** {ch['focus']}", "", f"**Dramatic question:** {ch['question']}", "", f"**Beginning:** {ch['start']}", "", f"**Ending:** {ch['end']}", ""]
        for act in ch["acts"]:
            part += [f"## {act['title']}", "", act["text"], "", f"**Turning point:** {act['turn']}", ""]
        part += ["## Possible short episodes", ""]
        for i, (title, text) in enumerate(ch["episodes"], 1):
            part += [f"{i}. **{title}:** {text}"]
        part += ["", "## Memory prompts", ""]
        for i, prompt in enumerate(ch["prompts"], 1):
            part += [f"{i}. {prompt}", "", "   Notes: ", ""]
        part += ["**Grand-Chukuchuku motif:** " + ch["device"], "", "**Details still open:** " + ch["continuity"], ""]
        filename = f"{ch['number']:02d}_{ch['slug']}.md"
        (ROOT / "editable" / filename).write_text("\n".join(part))
        lines += [""] + part
    lines += ["# Expansion guide", "", "For each memory, record approximate date, place, people, what Zino wanted, what stood in his way, what he did, and what changed. Add the food, music, speech, routines, and sensory details that make the scene feel lived.", "", "Distinguish personal recollection, information from someone else, a record, and proposed dramatization. Keep precise remembered words separate from reconstructed dialogue. Public history still requires verification.", "", "For each 1-3 minute episode: Act I establishes an immediate desire and problem; Act II complicates the attempt; Act III gives a consequence and a reason to continue. Pay off the preceding episode's question.", "", "The house remains unbuilt. Develop redemption from an actual decision, relationship, or change, without inventing a completed goal or a fully resolved emotional recovery.", ""]
    (ROOT / "editable" / "Grand-Chukuchuku_Editable_Outline.md").write_text("\n".join(lines))

def export_word():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[name].font.color.rgb = RGBColor.from_string("183C39")
    header = section.header.paragraphs[0]
    header.text = "GRAND-CHUKUCHUKU  |  EDITABLE DEVELOPMENT OUTLINE"
    header.style = doc.styles["Caption"]
    footer = section.footer.paragraphs[0]
    footer.text = "Based on Zino's account | Proposed structure for expansion | 28 October 2026"
    footer.style = doc.styles["Caption"]
    for line in (ROOT / "editable" / "Grand-Chukuchuku_Editable_Outline.md").read_text().splitlines():
        if not line.strip():
            continue
        if line.startswith("# "):
            text = line[2:]
            if text.startswith("Chapter ") or text == "Expansion guide":
                doc.add_page_break()
            doc.add_heading(text, 0 if text == "GRAND-CHUKUCHUKU" else 1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 2)
        else:
            paragraph = doc.add_paragraph()
            for i, fragment in enumerate(re.split(r"\*\*(.*?)\*\*", line)):
                run = paragraph.add_run(fragment)
                run.bold = i % 2 == 1
            if line.strip() == "Notes:":
                for _ in range(2):
                    doc.add_paragraph("________________________________________________________________")
    doc.core_properties.title = "Grand-Chukuchuku | Editable Series and Chapter Outline"
    doc.core_properties.author = "Based on the life of Nkeh Nkengzincho (Zino)"
    doc.save(ROOT / "editable" / "Grand-Chukuchuku_Editable_Outline.docx")

def main():
    ROOT.mkdir(parents=True, exist_ok=True)
    master = ROOT / "Grand-Chukuchuku_Series_and_Chapters.pdf"
    c = make_canvas(master)
    cover(c)
    foundation(c)
    synopsis(c)
    people(c)
    world(c)
    chronology(c)
    for ch in DATA["chapters"]:
        page = 7 + 2 * (ch["number"] - 1)
        chapter_outline(c, ch, page, 17)
        chapter_workbook(c, ch, page + 1, 17)
    expansion(c)
    c.save()
    for ch in DATA["chapters"]:
        target = ROOT / "chapters" / f"{ch['number']:02d}_{ch['slug']}.pdf"
        c = make_canvas(target)
        c.setTitle(f"Grand-Chukuchuku | Chapter {ch['number']}: {ch['title']}")
        chapter_outline(c, ch, 1, 2)
        chapter_workbook(c, ch, 2, 2)
        c.save()
    export_markdown()
    export_word()
    readme = """GRAND-CHUKUCHUKU - DEVELOPMENT PACK

Start with Grand-Chukuchuku_Series_and_Chapters.pdf (17 pages).
It contains the unified synopsis, characters, Fiango world, chronology, and all five chapters.

chapters/ contains five separate two-page PDFs. Each has three acts,
turning points, possible 1-3 minute episodes, and memory-expansion prompts.

editable/ contains an editable Word outline, the master Markdown outline,
five chapter Markdown files, structured source JSON, and the Python builder
(requires reportlab and python-docx).
The Markdown files can be opened in any text editor or imported into a writing app.
The structured JSON is the authoritative source for rebuilding the PDFs.

scripts/ contains episode-screenplay drafts. CONTINUITY_LOG.md records
confirmed corrections and questions that must be resolved before scripting.
writers-room/ contains the active season map, cultural rules, fact tiers, and
editorial requirements used to develop each new script.

All biographical material is based on Zino's account, not independently verified.
Proposed scenes and disputed details are identified for further development.
This is a story-development draft, not a finished screenplay or produced series.
"""
    (ROOT / "README.txt").write_text(readme)
    zip_path = ROOT / "Grand-Chukuchuku_Complete_Development_Pack.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        for p in [master, ROOT / "README.txt", ROOT / "CONTINUITY_LOG.md"] + sorted((ROOT / "chapters").glob("*.pdf")) + sorted((ROOT / "editable").glob("*")) + sorted((ROOT / "scripts").glob("*")) + sorted((ROOT / "writers-room").glob("*")):
            if p.is_file():
                z.write(p, p.relative_to(ROOT))
    print(json.dumps({"master": str(master), "archive": str(zip_path), "chapter_pdfs": 5}, indent=2))

if __name__ == "__main__":
    main()
