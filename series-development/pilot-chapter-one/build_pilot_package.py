#!/usr/bin/env python3
"""Build the five-minute Chapter One pilot screenplay and visual storyboard."""

from __future__ import annotations

from html import escape
from pathlib import Path
import re
import shutil
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parent
SERIES = ROOT.parent
PROJECT = SERIES.parent
SCRIPT_MD = ROOT / "Grand-Chukuchuku_Chapter_01_Pilot_The_Courtyard_Keeps_Score.md"
BOARD_MD = ROOT / "Grand-Chukuchuku_Chapter_01_Pilot_Visual_Storyboard.md"
SCRIPT_PDF = SCRIPT_MD.with_suffix(".pdf")
SCRIPT_DOCX = SCRIPT_MD.with_suffix(".docx")
BOARD_PDF = ROOT / "Grand-Chukuchuku_Chapter_01_Pilot_Visual_Storyboard.pdf"
PACKAGE = SERIES / "Grand-Chukuchuku_Chapter_01_5_Minute_Pilot_Package.zip"
OUTPUT = PROJECT.parent / "output" / "grand-chukuchuku"

INK = colors.HexColor("#183C39")
MUTED = colors.HexColor("#586863")
GOLD = colors.HexColor("#BD8D42")
PAPER = colors.HexColor("#FAF8F2")
PALE = colors.HexColor("#E8EEE7")
LINE = colors.HexColor("#D2DAD2")
W, H = A4
M = 42


def fonts() -> None:
    root = Path("/usr/share/fonts/truetype/dejavu")
    for name, filename in [
        ("PilotSans", "DejaVuSans.ttf"),
        ("PilotSansB", "DejaVuSans-Bold.ttf"),
        ("PilotSerif", "DejaVuSerif.ttf"),
        ("PilotSerifB", "DejaVuSerif-Bold.ttf"),
    ]:
        pdfmetrics.registerFont(TTFont(name, str(root / filename)))
    pdfmetrics.registerFontFamily("PilotSans", normal="PilotSans", bold="PilotSansB")


def clean(value: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", value).replace("`", "").strip()


def rich(value: str) -> str:
    safe = escape(value.strip())
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", safe)


def script_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("pilot-title", parent=base["Title"], fontName="PilotSerifB", fontSize=25, leading=30, textColor=INK, spaceAfter=10),
        "h1": ParagraphStyle("pilot-h1", parent=base["Heading1"], fontName="PilotSerifB", fontSize=17, leading=21, textColor=INK, spaceBefore=10, spaceAfter=8, keepWithNext=True),
        "h2": ParagraphStyle("pilot-h2", parent=base["Heading2"], fontName="PilotSansB", fontSize=11.8, leading=16, textColor=INK, spaceBefore=9, spaceAfter=5, keepWithNext=True),
        "body": ParagraphStyle("pilot-body", parent=base["BodyText"], fontName="PilotSans", fontSize=9.45, leading=13.4, textColor=INK, spaceAfter=6),
        "bullet": ParagraphStyle("pilot-bullet", parent=base["BodyText"], fontName="PilotSans", fontSize=9.2, leading=13.1, textColor=INK, leftIndent=15, firstLineIndent=-8, spaceAfter=4),
        "small": ParagraphStyle("pilot-small", parent=base["BodyText"], fontName="PilotSans", fontSize=8.25, leading=11.5, textColor=MUTED, spaceAfter=5),
    }


def chrome(c: canvas.Canvas, doc: SimpleDocTemplate) -> None:
    c.saveState()
    c.setFillColor(PAPER)
    c.rect(0, 0, W, H, fill=1, stroke=0)
    c.setStrokeColor(LINE)
    c.line(M, H - 39, W - M, H - 39)
    c.line(M, 42, W - M, 42)
    c.setFillColor(MUTED)
    c.setFont("PilotSansB", 7.3)
    c.drawString(M, H - 29, "GRAND-CHUKUCHUKU")
    c.setFont("PilotSans", 7.1)
    c.drawRightString(W - M, H - 29, "CHAPTER ONE PILOT · THE COURTYARD KEEPS SCORE")
    c.drawString(M, 29, "AUTHOR-BASED DEVELOPMENT DRAFT · 29 OCTOBER 2026")
    c.setFont("PilotSansB", 7.5)
    c.drawRightString(W - M, 29, f"{doc.page:02d}")
    c.restoreState()


def build_script_pdf() -> None:
    styles = script_styles()
    flow = []
    first_title = True
    for raw in SCRIPT_MD.read_text().splitlines():
        line = raw.strip()
        if not line or line == "---":
            flow.append(Spacer(1, 4))
        elif line.startswith("# "):
            flow.append(Paragraph(rich(clean(line[2:])), styles["title"] if first_title else styles["h1"]))
            first_title = False
        elif line.startswith("## "):
            heading = clean(line[3:])
            if heading == "Timed shooting script":
                flow.append(PageBreak())
            flow.append(Paragraph(rich(heading), styles["h1"] if heading == "Timed shooting script" else styles["h2"]))
        elif line.startswith("### "):
            flow.append(Paragraph(rich(clean(line[4:])), styles["h2"]))
        elif line.startswith("- "):
            flow.append(Paragraph("• " + rich(line[2:]), styles["bullet"]))
        elif re.match(r"^\d+\.\s", line):
            flow.append(Paragraph(rich(line), styles["bullet"]))
        elif line.startswith("|"):
            continue
        else:
            style = styles["small"] if line.startswith("**C —") or line.startswith("**P —") else styles["body"]
            flow.append(Paragraph(rich(line), style))
    doc = SimpleDocTemplate(str(SCRIPT_PDF), pagesize=A4, leftMargin=M, rightMargin=M, topMargin=57, bottomMargin=57,
                            title="Grand-Chukuchuku | The Courtyard Keeps Score", author="Based on the life of Nkeh Nkengzincho (Zino)")
    doc.build(flow, onFirstPage=chrome, onLaterPages=chrome)


def add_rich_docx(paragraph, value: str) -> None:
    for index, piece in enumerate(re.split(r"(\*\*.*?\*\*)", value)):
        if not piece:
            continue
        run = paragraph.add_run(clean(piece))
        run.bold = piece.startswith("**") and piece.endswith("**")


def build_script_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.7)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    for style in ["Title", "Heading 1", "Heading 2"]:
        doc.styles[style].font.color.rgb = RGBColor.from_string("183C39")
    header = section.header.paragraphs[0]
    header.text = "GRAND-CHUKUCHUKU  |  CHAPTER ONE PILOT"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    footer.text = "Author-based development draft · 29 October 2026"
    for raw in SCRIPT_MD.read_text().splitlines():
        line = raw.strip()
        if not line or line == "---" or line.startswith("|"):
            continue
        if line.startswith("# "):
            doc.add_heading(clean(line[2:]), 0)
        elif line.startswith("## "):
            if clean(line[3:]) == "Timed shooting script":
                doc.add_page_break()
            doc.add_heading(clean(line[3:]), 1)
        elif line.startswith("### "):
            doc.add_heading(clean(line[4:]), 2)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_rich_docx(p, line[2:])
        else:
            p = doc.add_paragraph()
            add_rich_docx(p, line)
    doc.core_properties.title = "Grand-Chukuchuku | The Courtyard Keeps Score"
    doc.core_properties.author = "Based on the life of Nkeh Nkengzincho (Zino)"
    doc.save(SCRIPT_DOCX)


def board_chrome(c: canvas.Canvas, page: int, pages: int) -> None:
    c.setFillColor(PAPER)
    c.rect(0, 0, W, H, fill=1, stroke=0)
    c.setStrokeColor(LINE)
    c.line(M, H - 39, W - M, H - 39)
    c.line(M, 42, W - M, 42)
    c.setFillColor(MUTED)
    c.setFont("PilotSansB", 7.3)
    c.drawString(M, H - 29, "GRAND-CHUKUCHUKU")
    c.setFont("PilotSans", 7.1)
    c.drawRightString(W - M, H - 29, "5-MINUTE PILOT · VISUAL STORYBOARD")
    c.drawString(M, 29, "AUTHOR-BASED DEVELOPMENT DRAFT · 29 OCTOBER 2026")
    c.setFont("PilotSansB", 7.5)
    c.drawRightString(W - M, 29, f"{page:02d} / {pages:02d}")


def image_path(name: str) -> Path:
    mapping = {
        "anchor": SERIES / "picture-board" / "00_Fiango_Character_Anchor.png",
        "bottle": SERIES / "picture-board" / "05_Nacking_Bottle.png",
        "school": SERIES / "picture-board" / "06_Between_Bells.png",
        "fantasy": ROOT / "images" / "01_Grand_Bends_the_Ball.png",
        "notice": ROOT / "images" / "02_The_Fee_Notice.png",
    }
    return mapping[name]


BOARDS = [
    ("01", "0:00–0:08", "anchor", "COLD OPEN: THE BALL", "Macro on a spinning, scuffed football moving through red dust. Rack focus to two slippers marking a narrow goal.", "Zino O.S. · ball thump, bucket rope, distant radio", "Hard cut on impact."),
    ("02", "0:08–0:21", "anchor", "THE COMPOUND", "Tilt from mango leaves to Zino, Gideon and Luc. Keep the well and close-set homes legible behind them.", "Gideon scores. Luc asks for fairness.", "Whip pan with Zino’s run."),
    ("03", "0:21–0:40", "anchor", "THE MISS", "Low angle follows Zino’s shot. Hold a half beat as it brushes the outer slipper; reveal the plastic basin.", "“Outside. No goal.”", "Zino looks for listeners."),
    ("04", "0:40–0:58", "fantasy", "THE FIRST LEGEND", "Match cut the ball’s arc into gold. Push toward Grand on the roofline; he bends the ball, but never flies.", "Zino’s first impossible story.", "Gold arc wipes into water ripple."),
    ("05", "0:58–1:18", "anchor", "BACK TO EARTH", "Real ball in basin. Reveal Mother, carrying marked exercise books. Her calm command changes the energy.", "“Basin first. Football later.”", "Cut on her pointed look."),
    ("06", "1:18–1:38", "anchor", "WELL WATER", "Side-on medium: Zino rushes the rope, gets splashed, tries again slowly. Gideon and Luc occupy separate thirds.", "Friendly teasing; rope squeal.", "Rope rises into notebook edge."),
    ("07", "1:38–1:58", "notice", "THE PRICE IN THE COURTYARD", "Long-lens, background action: landlord works door to door. Rack focus to Zino’s hand still on the well rope.", "“Rent and light matter.” Keep amount unheard.", "Notebook cover fills frame."),
    ("08", "1:58–2:16", "anchor", "A SMALL PLAN", "Mother puts a wrapped lunch and coins in Zino’s school bag. In the distance: bottles and a cloth of pooled stakes.", "“Your head is not your stomach.”", "Match cut coin to school desk."),
    ("09", "2:16–2:38", "school", "SCHOOL MORNING", "Wide lunch exodus, then isolate Zino as he folds coins into an exercise-book page. Gideon and Luc notice.", "“Hunger no be strength.”", "Bell tail becomes chalk scratch."),
    ("10", "2:38–3:00", "school", "THE ANSWER", "Push from the exercise book to Zino’s face as he answers correctly. Snap down to catch the falling folded paper.", "Teacher: “Correct.”", "Paper becomes white cloth under coins."),
    ("11", "3:00–3:18", "bottle", "THE OFFER", "Ground-level insert of bottles, slippers and small pooled coins. Rise to reveal the street circle around the boys.", "On-screen rule: hit opposing bottle three times.", "Ball rolls at lens."),
    ("12", "3:18–3:37", "bottle", "FIRST HIT", "Slow anticipation, then normal speed when the ball clips the bottle. Track its fall to one raised finger.", "Compact cheer; Gideon grins.", "Cut on glass clink."),
    ("13", "3:37–3:58", "bottle", "THE TEST", "Reverse field geography as the opponent strikes Zino’s bottle. Move from coins into Zino’s changing expression.", "Luc checks his temper.", "Faint gold edge at bottle neck."),
    ("14", "3:58–4:16", "fantasy", "THE SHOT HE CANNOT BUY", "Grand remains a faint outline and does not help. Zino misses; hold dust travelling past the standing bottle.", "Mother calls from the compound.", "Her call bridges us home."),
    ("15", "4:16–4:34", "notice", "THE NOTE", "Low follow as Zino runs home. Mother waits in doorway with school notice and books; hide his coin fist.", "“Come inside.”", "Doorway shadow takes frame."),
    ("16", "4:34–4:48", "notice", "THE QUESTION", "Over Zino’s shoulder. Keep the letter’s content unreadable. Push slowly toward his closed fist.", "“You know for what.”", "Strip score to room tone."),
    ("17", "4:48–5:00", "notice", "CLIFFHANGER", "Custom insert: Zino opens his hand. Coin shine becomes the fantasy’s gold ball arc; cut before the score.", "Zino V.O.; distant radio; title card.", "Black → Episode 02."),
]


def cover(c: canvas.Canvas) -> None:
    c.setFillColor(PAPER)
    c.rect(0, 0, W, H, fill=1, stroke=0)
    c.drawImage(str(image_path("fantasy")), 0, 0, width=W, height=H, preserveAspectRatio=True, anchor="c", mask="auto")
    c.setFillColor(colors.Color(0.06, 0.12, 0.12, alpha=.72))
    c.rect(0, 0, W, 210, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#F9F2E4"))
    c.setFont("PilotSansB", 9)
    c.drawString(M, 174, "CHAPTER ONE · 5-MINUTE PILOT")
    c.setFont("PilotSerifB", 31)
    c.drawString(M, 126, "THE COURTYARD")
    c.drawString(M, 86, "KEEPS SCORE")
    c.setFont("PilotSans", 10)
    c.drawString(M, 60, "Visual storyboard + timed shooting script")
    c.setFillColor(GOLD)
    c.rect(M, 188, 72, 4, fill=1, stroke=0)
    c.showPage()


def fit_image(c: canvas.Canvas, image: Path, x: float, y: float, w: float, h: float) -> None:
    with PILImage.open(image) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    dw, dh = iw * scale, ih * scale
    c.setFillColor(colors.white)
    c.roundRect(x, y, w, h, 5, fill=1, stroke=0)
    c.drawImage(str(image), x + (w - dw) / 2, y + (h - dh) / 2, width=dw, height=dh, mask="auto")


def board_card(c: canvas.Canvas, row, top: float) -> None:
    number, timing, source, title, framing, sound, transition = row
    card_h = 325
    x = M
    c.setFillColor(colors.white)
    c.roundRect(x, top-card_h, W-2*M, card_h, 8, fill=1, stroke=0)
    c.setStrokeColor(LINE)
    c.roundRect(x, top-card_h, W-2*M, card_h, 8, fill=0, stroke=1)
    image_w, image_h = 150, 277
    fit_image(c, image_path(source), x+13, top-card_h+24, image_w, image_h)
    tx = x + 180
    c.setFillColor(GOLD)
    c.setFont("PilotSansB", 8)
    c.drawString(tx, top-25, f"BOARD {number}  ·  {timing}")
    c.setFillColor(INK)
    c.setFont("PilotSerifB", 15)
    c.drawString(tx, top-47, title)
    text_w = W - M - tx - 14
    styles = {
        "body": ParagraphStyle("board-body", fontName="PilotSans", fontSize=8.7, leading=12, textColor=INK),
        "small": ParagraphStyle("board-small", fontName="PilotSans", fontSize=8.1, leading=11.2, textColor=MUTED),
    }
    blocks = [("CAMERA / ACTION", framing), ("SOUND / PERFORMANCE", sound), ("CUT", transition)]
    yy = top-71
    for label, text in blocks:
        c.setFillColor(GOLD)
        c.setFont("PilotSansB", 7)
        c.drawString(tx, yy, label)
        yy -= 5
        para = Paragraph(escape(text), styles["body"] if label == "CAMERA / ACTION" else styles["small"])
        _, height = para.wrap(text_w, 200)
        para.drawOn(c, tx, yy-height)
        yy -= height + 13


def build_board_pdf() -> None:
    total_pages = 10
    c = canvas.Canvas(str(BOARD_PDF), pagesize=A4, pageCompression=1)
    c.setTitle("Grand-Chukuchuku | The Courtyard Keeps Score — Visual Storyboard")
    c.setAuthor("Based on the life of Nkeh Nkengzincho (Zino)")
    cover(c)
    for page_index in range(9):
        board_chrome(c, page_index+2, total_pages)
        start = page_index * 2
        rows = BOARDS[start:start+2]
        y = H - 62
        for row in rows:
            board_card(c, row, y)
            y -= 346
        c.showPage()
    c.save()


def write_readme() -> Path:
    target = ROOT / "README.md"
    target.write_text("""# Chapter One pilot package\n\n**The Courtyard Keeps Score** is a five-minute vertical opening pilot for *Grand-Chukuchuku*. It introduces Fiango, Zino, Gideon, Luc, Zino’s mother, the imaginary adult Grand-Chukuchuku, Nacking Bottle, and primary-school fee pressure.\n\nStart with `Grand-Chukuchuku_Chapter_01_Pilot_Visual_Storyboard.pdf`, then use the script PDF or editable DOCX for the voice recording and animatic. The storyboard has 17 timed boards and runs exactly 5:00.\n\nThe screen action, dialogue, fee notice and exact chronology are proposed dramatic construction. The accompanying script marks the confirmed factual foundations and the safeguards that prevent this pilot from claiming a literal reconstruction.\n\nThe two images in `images/` are newly generated visual-development keyframes. The picture-board panels in `Reference_Assets/` are existing Chapter One visual development assets. They are all subject to the author-supplied baseline in `../visual-reference/VISUAL_BASELINE_V1.md`.\n""")
    return target


def build_package() -> None:
    readme = write_readme()
    files = [SCRIPT_MD, SCRIPT_PDF, SCRIPT_DOCX, BOARD_MD, BOARD_PDF, readme,
             ROOT / "images" / "01_Grand_Bends_the_Ball.png", ROOT / "images" / "02_The_Fee_Notice.png",
             SERIES / "picture-board" / "00_Fiango_Character_Anchor.png",
             SERIES / "picture-board" / "05_Nacking_Bottle.png",
             SERIES / "picture-board" / "06_Between_Bells.png",
             SERIES / "visual-reference" / "VISUAL_BASELINE_V1.md"]
    with zipfile.ZipFile(PACKAGE, "w", zipfile.ZIP_DEFLATED) as z:
        for file in files:
            if file.parent == ROOT / "images":
                arcname = Path("images") / file.name
            elif file.parent == SERIES / "picture-board":
                arcname = Path("Reference_Assets") / file.name
            elif file.parent == SERIES / "visual-reference":
                arcname = Path("Reference_Assets") / file.name
            else:
                arcname = file.name
            z.write(file, arcname)
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for file in [PACKAGE, SCRIPT_PDF, BOARD_PDF]:
        shutil.copy2(file, OUTPUT / file.name)


def main() -> None:
    fonts()
    build_script_pdf()
    build_script_docx()
    build_board_pdf()
    build_package()
    print("Built:")
    for file in [SCRIPT_PDF, SCRIPT_DOCX, BOARD_PDF, PACKAGE]:
        print(file)


if __name__ == "__main__":
    main()
