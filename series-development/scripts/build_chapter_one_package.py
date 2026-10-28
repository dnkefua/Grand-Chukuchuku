#!/usr/bin/env python3
"""Build the Chapter One shooting-script and illustrated picture-board deliverables."""

from __future__ import annotations

from html import escape
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parent.parent
SCRIPT = ROOT / "scripts" / "Grand-Chukuchuku_Chapter_01_Full_Shooting_Script_v1.md"
SCRIPT_PDF = SCRIPT.with_suffix(".pdf")
SCRIPT_DOCX = SCRIPT.with_suffix(".docx")
BOARD_DIR = ROOT / "picture-board"
BOARD_PDF = BOARD_DIR / "Grand-Chukuchuku_Chapter_01_Picture_Board.pdf"
PACKAGE = ROOT / "Grand-Chukuchuku_Chapter_01_Script_and_Picture_Board_v1.zip"
CONTINUITY = ROOT / "CONTINUITY_LOG.md"
REVIEW = ROOT / "writers-room" / "drafts" / "CRITIC_CHAPTER_01_REVIEW.md"

INK = colors.HexColor("#183C39")
MUTED = colors.HexColor("#586863")
GOLD = colors.HexColor("#BD8D42")
PAPER = colors.HexColor("#FAF8F2")
PALE = colors.HexColor("#E8EEE7")
LINE = colors.HexColor("#D2DAD2")
W, H = A4

PANELS = [
    (
        "00_Fiango_Character_Anchor.png",
        "Fiango: the character anchor",
        "Episodes 01–06 · Zino, Gideon, Luc, Zino’s mother, football, the well and the mango tree establish the world before its pressure closes in.",
    ),
    (
        "01_Professional_School_Night.png",
        "Professional School at night",
        "Episode 07 · A real vocational campus by day becomes a shared revision space at night: focus, temptation and escape in the same corridor.",
    ),
    (
        "02_Radio_Results.png",
        "Results on air",
        "Episode 09 · The Ordinary-Level radio broadcast turns the quiet wait into public release and opens the boarding-school choice.",
    ),
    (
        "03_The_Long_Road.png",
        "The long road",
        "Episodes 10–12 · A day through mountains and valleys carries Zino from Fiango into boarding-school discipline and a different social world.",
    ),
    (
        "04_The_Crowd_Witness_View.png",
        "The crowd: witness view",
        "Episode 08 · A non-spectacular witness image for the hard-R sequence. The accused remains unnamed and guilt is never established.",
    ),
]


def register_fonts() -> None:
    root = Path("/usr/share/fonts/truetype/dejavu")
    for name, filename in [
        ("GCSans", "DejaVuSans.ttf"),
        ("GCSansB", "DejaVuSans-Bold.ttf"),
        ("GCSerif", "DejaVuSerif.ttf"),
        ("GCSerifB", "DejaVuSerif-Bold.ttf"),
    ]:
        pdfmetrics.registerFont(TTFont(name, str(root / filename)))
    pdfmetrics.registerFontFamily("GCSans", normal="GCSans", bold="GCSansB")


def markdown_rich(value: str) -> str:
    safe = escape(value.strip())
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", safe)


def clean_markdown(value: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", value).replace("`", "").strip()


def styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("gc-title", parent=base["Title"], fontName="GCSerifB", fontSize=25, leading=30, textColor=INK, spaceAfter=12),
        "subtitle": ParagraphStyle("gc-subtitle", parent=base["Normal"], fontName="GCSans", fontSize=11, leading=16, textColor=MUTED, spaceAfter=15),
        "h1": ParagraphStyle("gc-h1", parent=base["Heading1"], fontName="GCSerifB", fontSize=17, leading=21, textColor=INK, spaceBefore=8, spaceAfter=9, keepWithNext=True),
        "h2": ParagraphStyle("gc-h2", parent=base["Heading2"], fontName="GCSansB", fontSize=12.2, leading=16, textColor=INK, spaceBefore=9, spaceAfter=5, keepWithNext=True),
        "h3": ParagraphStyle("gc-h3", parent=base["Heading3"], fontName="GCSansB", fontSize=10.5, leading=14, textColor=GOLD, spaceBefore=8, spaceAfter=4, keepWithNext=True),
        "body": ParagraphStyle("gc-body", parent=base["BodyText"], fontName="GCSans", fontSize=9.35, leading=13.3, textColor=INK, spaceAfter=6),
        "small": ParagraphStyle("gc-small", parent=base["BodyText"], fontName="GCSans", fontSize=8.2, leading=11.4, textColor=MUTED, spaceAfter=5),
        "bullet": ParagraphStyle("gc-bullet", parent=base["BodyText"], fontName="GCSans", fontSize=9.15, leading=13.1, textColor=INK, leftIndent=14, firstLineIndent=-8, spaceAfter=4),
    }


def page_chrome(c: canvas.Canvas, doc: SimpleDocTemplate) -> None:
    c.saveState()
    c.setFillColor(PAPER)
    c.rect(0, 0, W, H, fill=1, stroke=0)
    c.setStrokeColor(LINE)
    c.line(43, H - 39, W - 43, H - 39)
    c.line(43, 43, W - 43, 43)
    c.setFillColor(MUTED)
    c.setFont("GCSansB", 7.2)
    c.drawString(43, H - 29, "GRAND-CHUKUCHUKU")
    c.setFont("GCSans", 7.1)
    c.drawRightString(W - 43, H - 29, "CHAPTER ONE · FIANGO · SHOOTING DRAFT")
    c.drawString(43, 29, "AUTHOR-BASED DEVELOPMENT DRAFT · 28 OCTOBER 2026")
    c.setFont("GCSansB", 7.5)
    c.drawRightString(W - 43, 29, f"{doc.page:02d}")
    c.restoreState()


def script_flowables(markdown: str) -> list:
    s = styles()
    out: list = []
    lines = markdown.splitlines()
    i = 0
    first_title = True
    intro_arc = False
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line or line == "---":
            out.append(Spacer(1, 4))
            i += 1
            continue
        if line.startswith("|") and line.endswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [clean_markdown(c.strip()) for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[: -]+", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                parsed = [[Paragraph(markdown_rich(c), s["small"] if r else s["body"]) for c in row] for r, row in enumerate(rows)]
                widths = [(W - 86) / len(rows[0])] * len(rows[0])
                tbl = Table(parsed, colWidths=widths, repeatRows=1)
                tbl.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), PALE),
                    ("GRID", (0, 0), (-1, -1), 0.35, LINE),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]))
                out.extend([tbl, Spacer(1, 7)])
            continue
        if line.startswith("# "):
            heading = clean_markdown(line[2:])
            if not first_title and "EPISODE" in heading.upper():
                out.append(PageBreak())
                intro_arc = False
            out.append(Paragraph(markdown_rich(heading), s["title"] if first_title else s["h1"]))
            first_title = False
        elif line.startswith("## "):
            heading = clean_markdown(line[3:])
            if "EPISODE" in heading.upper():
                out.append(PageBreak())
                out.append(Paragraph(markdown_rich(heading), s["h1"]))
                intro_arc = False
            else:
                out.append(Paragraph(markdown_rich(heading), s["h2"]))
                intro_arc = heading.lower() == "one-page chapter arc"
        elif line.startswith("### "):
            out.append(Paragraph(markdown_rich(clean_markdown(line[4:])), s["h3"]))
        elif line.startswith("- "):
            out.append(Paragraph("• " + markdown_rich(line[2:]), s["bullet"]))
        elif re.match(r"^\d+\.\s", line):
            out.append(Paragraph(markdown_rich(line), s["bullet"]))
        else:
            style = s["small"] if intro_arc or line.startswith("**Episode fact boundary") or line.startswith("**Fact boundary") else s["body"]
            out.append(Paragraph(markdown_rich(line), style))
        i += 1
    return out


def build_script_pdf(markdown: str) -> None:
    doc = SimpleDocTemplate(
        str(SCRIPT_PDF), pagesize=A4, leftMargin=43, rightMargin=43,
        topMargin=57, bottomMargin=57, title="Grand-Chukuchuku | Chapter One Full Shooting Script",
        author="Based on the life of Nkeh Nkengzincho (Zino)",
    )
    doc.build(script_flowables(markdown), onFirstPage=page_chrome, onLaterPages=page_chrome)


def word_add_rich(paragraph, value: str) -> None:
    for index, part in enumerate(re.split(r"(\*\*.*?\*\*)", value)):
        if not part:
            continue
        run = paragraph.add_run(clean_markdown(part))
        run.bold = part.startswith("**") and part.endswith("**")


def build_script_docx(markdown: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("183C39")
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].font.color.rgb = RGBColor.from_string("183C39")
    header = section.header.paragraphs[0]
    header.text = "GRAND-CHUKUCHUKU  |  CHAPTER ONE · FIANGO · SHOOTING DRAFT"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer = section.footer.paragraphs[0]
    footer.text = "Author-based development draft · 28 October 2026"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first = True
    lines = markdown.splitlines()
    for raw in lines:
        line = raw.strip()
        if not line or line == "---":
            continue
        if line.startswith("# "):
            text = clean_markdown(line[2:])
            if not first and "EPISODE" in text.upper():
                doc.add_page_break()
            doc.add_heading(text, 0 if first else 1)
            first = False
        elif line.startswith("## "):
            text = clean_markdown(line[3:])
            if "EPISODE" in text.upper():
                doc.add_page_break()
                doc.add_heading(text, 1)
            else:
                doc.add_heading(text, 2)
        elif line.startswith("### "):
            doc.add_heading(clean_markdown(line[4:]), 3)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            word_add_rich(p, line[2:])
        else:
            p = doc.add_paragraph()
            word_add_rich(p, line)
    doc.core_properties.title = "Grand-Chukuchuku | Chapter One Full Shooting Script"
    doc.core_properties.author = "Based on the life of Nkeh Nkengzincho (Zino)"
    doc.core_properties.subject = "Fiango · 12 vertical animated episodes"
    doc.save(SCRIPT_DOCX)


def board_cover(c: canvas.Canvas) -> None:
    c.setFillColor(PAPER)
    c.rect(0, 0, W, H, fill=1, stroke=0)
    c.setFillColor(GOLD)
    c.rect(48, H - 95, 50, 5, fill=1, stroke=0)
    c.setFillColor(MUTED)
    c.setFont("GCSansB", 9)
    c.drawString(48, H - 73, "AN ILLUSTRATED DEVELOPMENT BOARD")
    c.setFillColor(INK)
    c.setFont("GCSerifB", 34)
    c.drawString(48, H - 146, "GRAND-")
    c.drawString(48, H - 188, "CHUKUCHUKU")
    c.setFont("GCSans", 14)
    c.drawString(48, H - 226, "Chapter One: Fiango")
    quote = Paragraph(
        "Five visual anchors for a 12-episode, 9:16 animated chapter about imagination, school, public violence, and the road away from home.",
        ParagraphStyle("quote", fontName="GCSerif", fontSize=14, leading=20, textColor=INK),
    )
    quote.wrapOn(c, W - 130, 150)
    quote.drawOn(c, 48, H - 355)
    c.setStrokeColor(LINE)
    c.line(48, 115, W - 48, 115)
    c.setFillColor(MUTED)
    c.setFont("GCSans", 8.5)
    c.drawString(48, 92, "AI-generated development art. Composite scenes; not literal documentary reconstruction.")
    c.drawString(48, 70, "Hard-R master requires a separate platform release cut before publishing to TikTok.")
    c.showPage()


def build_board_pdf() -> None:
    missing = [name for name, _, _ in PANELS if not (BOARD_DIR / name).exists()]
    if missing:
        raise FileNotFoundError("Missing board panels: " + ", ".join(missing))
    c = canvas.Canvas(str(BOARD_PDF), pagesize=A4, pageCompression=1)
    c.setTitle("Grand-Chukuchuku | Chapter One Picture Board")
    c.setAuthor("Based on the life of Nkeh Nkengzincho (Zino)")
    board_cover(c)
    for idx, (name, title, caption) in enumerate(PANELS, 1):
        c.setFillColor(PAPER)
        c.rect(0, 0, W, H, fill=1, stroke=0)
        c.setFillColor(GOLD)
        c.setFont("GCSansB", 8)
        c.drawString(43, H - 39, f"PANEL {idx:02d} / {len(PANELS):02d}")
        c.setFillColor(INK)
        c.setFont("GCSerifB", 19)
        c.drawString(43, H - 68, title)
        source = BOARD_DIR / name
        with PILImage.open(source) as im:
            iw, ih = im.size
        max_w, max_h = W - 86, H - 210
        scale = min(max_w / iw, max_h / ih)
        dw, dh = iw * scale, ih * scale
        x = (W - dw) / 2
        y = 112 + (max_h - dh) / 2
        c.drawImage(str(source), x, y, width=dw, height=dh, preserveAspectRatio=True, mask="auto")
        c.setStrokeColor(LINE)
        c.rect(x, y, dw, dh, stroke=1, fill=0)
        para = Paragraph(markdown_rich(caption), ParagraphStyle("cap", fontName="GCSans", fontSize=9.2, leading=13, textColor=MUTED))
        para.wrapOn(c, W - 86, 52)
        para.drawOn(c, 43, 62)
        c.setFillColor(MUTED)
        c.setFont("GCSans", 7)
        c.drawRightString(W - 43, 30, "CHAPTER ONE · PICTURE BOARD")
        c.showPage()
    c.save()


def make_package() -> None:
    targets = [
        SCRIPT, SCRIPT_PDF, SCRIPT_DOCX, BOARD_PDF, BOARD_DIR / "README.md",
        CONTINUITY, REVIEW,
        *[BOARD_DIR / name for name, _, _ in PANELS],
    ]
    with zipfile.ZipFile(PACKAGE, "w", zipfile.ZIP_DEFLATED) as z:
        for path in targets:
            if path == SCRIPT:
                dest = Path("Full_Shooting_Script") / path.name
            elif path in (SCRIPT_PDF, SCRIPT_DOCX):
                dest = Path("Full_Shooting_Script") / path.name
            elif path.parent == BOARD_DIR:
                dest = Path("Picture_Board") / path.name
            elif path == CONTINUITY:
                dest = Path("Documentation") / path.name
            else:
                dest = Path("Documentation") / path.name
            z.write(path, dest)


def main() -> None:
    register_fonts()
    if not SCRIPT.exists():
        raise FileNotFoundError(SCRIPT)
    markdown = SCRIPT.read_text(encoding="utf-8")
    build_script_pdf(markdown)
    build_script_docx(markdown)
    build_board_pdf()
    make_package()
    print("Built:")
    for path in (SCRIPT_PDF, SCRIPT_DOCX, BOARD_PDF, PACKAGE):
        print(path)


if __name__ == "__main__":
    main()
